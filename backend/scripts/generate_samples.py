import os
import sys
import docx
import asyncio
from datetime import datetime
from pymongo import MongoClient
from sentence_transformers import SentenceTransformer

# Adjust python path to allow importing app modules
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Imports from app
from app.config import settings
from app.utils.text_processor import TextProcessor
from app.vectorstore.store import vector_store

DOCS_DIR = "./sample_documents"
os.makedirs(DOCS_DIR, exist_ok=True)

# Define 52 distinct high-quality SaaS documents to generate
DOCUMENT_TOPICS = [
    # SaaS Enterprise Security Policies (1-10)
    ("SaaS_Security_Overview.txt", "SaaS Enterprise Security Overview", 
     "This document outlines the security architecture for the SaaS Platform. We employ AES-256 encryption at rest and TLS 1.3 in transit. Authentication is managed via Okta SSO with mandatory multi-factor authentication (MFA). Security audits are conducted bi-annually by external Certified Information Systems Auditors (CISA). Network segmentation is enforced across VPCs."),
    ("Database_Credentials_Encryption.txt", "Database Credentials Encryption Protocol",
     "Database credentials must never be stored in plain text or hardcoded in repositories. All database credentials, API keys, and private certificates must be stored in HashiCorp Vault. Vault access tokens are rotated every 30 days. Databases use IAM-based role authentication where supported, ensuring short-lived access."),
    ("Password_Rotation_Policy.docx", "Corporate Password Rotation Guidelines",
     "Enterprise passwords must meet the following complexity rules: minimum 14 characters, at least one uppercase letter, one lowercase letter, one number, and one special character. Passwords must be rotated every 90 days. Users cannot reuse their previous 5 passwords. Inactive accounts are suspended after 45 days of idle state."),
    ("Admin_Access_Provisioning.docx", "Administrator Access Provisioning SOP",
     "Requesting administrative privileges requires submitting an Access Request Ticket (ART) signed by the department head. Administrator accounts must use distinct 'admin-' prefixed emails. Access is reviewed quarterly. Temporary elevated privileges (just-in-time access) are limited to 4-hour windows via AWS IAM Identity Center."),
    ("Vulnerability_Management.txt", "System Vulnerability Management Standard",
     "Weekly automated vulnerability scans are run across all server instances and container registries using AWS Inspector and Prisma Cloud. High-severity vulnerabilities (CVSS score >= 7.0) must be patched within 7 business days. Critical patches (CVSS >= 9.0) must be applied within 48 hours. Vulnerability logs are archived in S3."),
    ("Data_Retention_Policy.docx", "Global Data Retention and Disposal Policy",
     "Customer database transactional records are retained for 7 years from transaction date to comply with financial regulations. Session logs are retained for 90 days. User account profiles are deleted 30 days after subscription termination. Data scrubbing uses cryptographic deletion (shredding) on block storage devices."),
    ("API_Rate_Limiting.txt", "External API Rate Limiting Specifications",
     "The API gateway enforces rate limits to prevent denial of service attacks. Standard tiers are limited to 1,000 requests per hour. Enterprise tiers enjoy 10,000 requests per hour. Rate limits are tracked using a Redis token bucket algorithm. Requests exceeding limits receive an HTTP 429 Too Many Requests status response."),
    ("VPC_Peering_Security.txt", "VPC Peering Security Controls",
     "VPC peering connections are restricted to internal microservice clusters. Inter-VPC traffic is routed strictly through AWS Transit Gateway with Security Groups locking port ranges. Peer networks must not have overlapping CIDR blocks. Port 22 (SSH) and 3389 (RDP) are disabled across VPC boundaries, relying on SSM Session Manager."),
    ("Multi-Tenant_Data_Isolation.docx", "Multi-Tenant Database Data Isolation Protocol",
     "Tenant isolation is maintained using logical schema separation (row-level security) in PostgreSQL databases. Every database query must append a tenant_id filter. Tenant context is populated via JWT tokens at the API middleware layer. Cross-tenant database queries are blocked by strict DB constraint policies."),
    ("Data_Privacy_Compliance.txt", "Global Data Privacy and GDPR Compliance Directive",
     "To maintain GDPR compliance, users possess the right to be forgotten. Upon receiving a privacy request, the compliance team initiates a script purging user identifiers from MongoDB, Elasticsearch, and Redis cache. Encrypted backups retain data for 30 days before natural rotation wipes records. PII is encrypted at rest."),

    # CI/CD & Devops Workflows (11-20)
    ("Staging_Deployment_Pipeline.docx", "Staging Environment Deployment Workflows",
     "The staging environment deployment runs automatically on merges to the 'release' branch. Steps include: 1. Code linting and static analysis (SonarQube). 2. Unit tests execution. 3. Docker image creation. 4. Deployment to EKS staging cluster. 5. E2E integration test suite execution. Staging metrics are logged in Prometheus."),
    ("Production_Release_Approvals.txt", "Production Release Approval Standards",
     "Deployments to production require a pull request approved by at least two senior engineers and a signed Change Advisory Board (CAB) ticket. Production releases are scheduled for Tuesdays and Thursdays between 05:00 and 07:00 UTC. Hotfixes require immediate VP of Engineering approval and automated smoke tests."),
    ("Docker_Build_Logging.txt", "Docker Build and Container Logging Setup",
     "All Docker container builds must export logs in JSON format to stdout. Container host nodes run Vector agents that forward stdout logs directly to Datadog. Docker files must use official base images from verified registries. Multi-stage builds are required to keep final container images under 150MB in size."),
    ("Kubernetes_Ingress_Controller.docx", "Kubernetes Ingress Controller Routing",
     "Production routing uses NGINX Ingress Controller. Ingress rules terminate TLS using certificates auto-renewed by Let's Encrypt via cert-manager. Path-based routing redirects traffic: /api/auth goes to auth-service; /api/chatbot goes to chatbot-service. We utilize ingress-based rate limit annotations to prevent spam."),
    ("Blue-Green_Deployment.txt", "Blue-Green Deployment Release Procedures",
     "We execute Blue-Green deployments for all stateless microservices. The router (Route53) splits traffic 90/10 during the first 10 minutes of active validation. Once smoke tests succeed on the green pool, 100% traffic is redirected. The previous blue pool remains standby for 60 minutes for immediate rollback ability."),
    ("Jenkins_Build_Agent.docx", "Jenkins Build Agent Configuration Guidelines",
     "Jenkins build nodes run on EC2 spot instances orchestrated by an Auto Scaling Group. Agents mount a shared EFS volume for workspace caching. Docker-in-Docker (DinD) is enabled on build agents using secured unix sockets. Build workspaces are automatically pruned of artifacts older than 5 days to preserve disk space."),
    ("Terraform_State_Management.txt", "Terraform State Locking and Storage SOP",
     "Terraform state files must be stored in a centralized, encrypted S3 bucket. State locking is enforced using a DynamoDB table. State buckets have Versioning enabled. Manual state editing is strictly forbidden. Changes to infrastructure must be committed to git and executed via GitHub Actions CI pipelines."),
    ("Chaos_Engineering_Testing.txt", "Chaos Engineering Resiliency Tests",
     "We run automated chaos testing weekly using Chaos Mesh in the staging environment. Tests inject latency into DB queries, kill pods randomly, and simulate network partitions. Microservices must demonstrate graceful degradation, responding with cached data or custom fallbacks. Target recovery time objective (RTO) is 15 seconds."),
    ("Docker_Image_Scanning.docx", "Container Security and Docker Image Scanning",
     "Trivy scans are integrated into the GitHub Actions workflow. Pull requests fail if a new container image contains vulnerabilities classified as CRITICAL or HIGH with an available patch. Base images must run as non-root users (UID 10001). We enforce distroless base images for production deployments."),
    ("Serverless_Functions_CICD.txt", "Serverless Lambda Function CI/CD Procedures",
     "AWS Lambda deployments are managed via Serverless Framework. CI/CD pipelines package functions, compile dependencies, and run unit tests. We enforce a 250MB size limit on zip deployments. Lambda functions are deployed using canary configurations (10% linear increments every 2 minutes) with automated rollbacks on CloudWatch alarms."),

    # Database Infrastructure & Cache (21-30)
    ("Database_Indexing_Optimization.docx", "Database Indexing and Query Optimizations",
     "To optimize database read operations, compound indexes must be built on high-cardinality keys. In MongoDB, create compound indexes matching query patterns. Avoid index explosion (maximum 64 indexes per collection). Slow queries taking > 100ms are logged in the slow-query log and reviewed weekly by DBAs."),
    ("Redis_Caching_Strategy.txt", "Redis Cache Eviction and TTL Standards",
     "Redis cache instances are configured with maxmemory-policy volatile-lru. All database cache entries must have a designated Time-To-Live (TTL). User sessions have a TTL of 2 hours. Config parameters and tenant settings have a TTL of 24 hours. Cache key formatting uses colon spacing: tenant:{id}:user:{id}."),
    ("PostgreSQL_Replication.docx", "PostgreSQL Replication and Failover SOP",
     "Production PostgreSQL uses a primary-replica model with synchronous streaming replication. Write queries hit the primary node, while reads are distributed across two read-replicas. Automated failover is managed by Patroni. Backup retention is managed via WAL-G, exporting WAL segments to S3 every 15 minutes."),
    ("MongoDB_Sharding_Guide.txt", "Enterprise MongoDB Sharding and Clustering",
     "The MongoDB database uses sharding for scale. The shard key is tenant_id, ensuring balanced distribution across three shard clusters. Config servers are deployed in a replica set across three availability zones. Balanced chunks are checked daily, and jumbo chunks are manually split during off-peak hours."),
    ("Elasticsearch_Indexing.docx", "Elasticsearch Cluster Index Management",
     "Elasticsearch handles logs and full-text search. Index lifecycle management (ILM) automatically rolls over indices exceeding 50GB or 30 days. Hot index nodes use NVMe SSD storage; warm index nodes utilize standard HDD volumes. Index templates define dynamic mappings, ensuring text fields have keyword sub-fields."),
    ("Database_Backup_Validation.txt", "Database Backup Restore Validation Protocol",
     "Automated daily database backups are stored in AWS S3 with Glacier vault locks. Every Sunday at 01:00 UTC, an automated restore test launches a temporary DB instance in an isolated subnet, restores the backup, and runs verification SQL tests. Restore status is sent to Slack. Backup failure immediately alerts on-call DBAs."),
    ("Connection_Pooling_Settings.txt", "HikariCP and PgBouncer Connection Pooling",
     "To avoid database thread exhaustion, microservices connecting to PostgreSQL must route traffic through PgBouncer connection poolers. Client connection pool size is capped at 20 active connections per container. Timeout settings: connection timeout 30s, idle timeout 10 minutes, max lifetime 30 minutes."),
    ("NoSQL_Migration_Workflows.docx", "NoSQL Schema Migration and Update SOP",
     "Schema updates in MongoDB must be executed via backward-compatible migrations. Scripts must handle default values for missing keys dynamically in the application models. Migrations are run using migrations packages during maintenance windows. Rolback scripts must be verified in staging before production runs."),
    ("Database_Monitoring_Metrics.txt", "Database Performance and Monitoring Metrics",
     "We monitor the following database health metrics in Datadog: 1. CPU Utilization (alert if > 80% for 5 mins). 2. Freeable Memory. 3. Active Connections Count. 4. Disk Queue Depth. 5. Index Hit Rate (should remain > 99%). Disk space exhaustion triggers alerts at 85% utilization, auto-expanding AWS EBS volumes."),
    ("DynamoDB_Global_Tables.docx", "DynamoDB Global Tables Multi-Region Architecture",
     "For multi-region high availability, user configuration tables are deployed as DynamoDB Global Tables replicated between us-east-1 and eu-west-1. Read capacity is managed using auto-scaling with a base provisioning of 50 RCUs. Write capacity uses on-demand scaling to support burst campaigns."),

    # Customer & HR Guidelines (31-40)
    ("Customer_Onboarding_SOP.docx", "SaaS Customer Onboarding Standard Procedures",
     "Customer onboarding involves: 1. Workspace provisioning (automated via Stripe webhook). 2. Default database schema setup. 3. Verification email dispatch. 4. Interactive in-app product tour trigger. Account managers receive a notification if a customer hasn't completed setup within 48 hours of signup."),
    ("HR_Onboarding_Handbook.pdf", "Employee HR Onboarding Handbook",
     "Welcome to the team! All new hires must complete security training within their first week. Laptop setup is managed via Jamf, enforcing disk encryption and automatic screen locking after 5 minutes of idle state. Default tools include Slack, Jira, GitHub, and 1Password. Direct deposit forms must be completed in Workday."),
    ("Remote_Work_Guidelines.txt", "Remote Work Security and Equipment Rules",
     "Employees are permitted to work remotely. All corporate network traffic must route through the corporate tailscale VPN. Working on public Wi-Fi without active VPN is strictly prohibited. Hardware keys (YubiKey) are required for signing into AWS and GitHub. Reimbursed home office setup includes a monitor and chair."),
    ("Client_SLA_Commitments.docx", "Client Service Level Agreement Commitments",
     "We commit to a 99.9% uptime SLA for enterprise clients, excluding planned maintenance windows. Planned maintenance is announced 5 business days in advance. Uptime is calculated monthly. Uptime infractions between 99.0% and 99.9% trigger a 10% billing credit; uptime below 99.0% triggers a 25% refund."),
    ("Incident_Response_Plan.pdf", "Cybersecurity Incident Response Framework",
     "In the event of a security breach, the following plan is executed: Phase 1: Identification of breach. Phase 2: Containment (revoke credentials, isolate subnets). Phase 3: Eradication of threat. Phase 4: Recovery (restore validated backups). Phase 5: Post-incident analysis. Incident reports must be completed in 72 hours."),
    ("Customer_Support_Triage.txt", "Customer Support Ticket Triage Guidelines",
     "Support tickets are categorized: Tier 1: General inquiries (response in 24 hours). Tier 2: Feature bug reports (response in 4 hours). Tier 3: Outage / Security incident (response in 15 minutes). Tier 3 tickets immediately trigger PagerDuty alerts to the on-call engineering team. Support uses Zendesk integrations."),
    ("Employee_Expense_Policy.docx", "Employee Expense Reimbursement Policy",
     "Expense reports must be submitted via Expensify within 30 days of purchase. Standard meals are capped at $50 per person. Travel airfares must be booked in economy class. Purchases of software licenses or hardware accessories exceeding $100 require pre-approval from the IT manager. Receipts are mandatory."),
    ("Code_of_Conduct.txt", "Company Code of Conduct and Anti-Harassment",
     "We foster an inclusive, collaborative, and diverse workspace. Harassment of any kind (verbal, physical, or online) is strictly prohibited. Violations lead to immediate termination. Disputes should be reported to the HR team. Confidentiality of internal documents and customer data must be maintained at all times."),
    ("Offboarding_Procedure.docx", "Employee Offboarding and Asset Recovery SOP",
     "Upon notification of departure, IT schedules automated account revoking at 17:00 local time on the employee's last day. Active accounts revoked: Google Workspace, Slack, AWS, GitHub, Okta, Jira. Company laptops must be returned to the office within 5 business days. Access badges are deactivated immediately."),
    ("Partner_Referral_Guidelines.txt", "Enterprise Partner Referral Guidelines",
     "Partners can refer prospective leads through registered agencies and external contractors. Referrals must be logged in the Partner Portal. Successful referrals that sign a 12-month contract earn a 15% commission on the first year's contract value. Commissions are paid out 30 days post invoice clearance."),

    # Software Engineering & Architecture (41-52)
    ("API_Design_Standards.docx", "RESTful API Design and Standards Guidelines",
     "API endpoints must follow REST design principles. Paths use plural resource naming (e.g., /api/documents). Response payloads must return JSON. Timestamps must conform to ISO 8601 formatting. Standard HTTP status codes: 200 OK, 201 Created, 400 Bad Request, 401 Unauthorized, 403 Forbidden, 404 Not Found, 500 Server Error."),
    ("Microservices_Communication.txt", "Microservice Inter-Service Communication Patterns",
     "Stateless microservices communicate asynchronously using RabbitMQ message brokers. Synchronous calls use gRPC for high performance. JSON over HTTP is restricted to external gateway requests. Message schemas must be documented in Confluence. Failed messages are routed to Dead Letter Exchanges (DLX) for analysis."),
    ("Frontend_State_Management.docx", "React Frontend State Management Architecture",
     "The React frontend manages global state using React Context API for themes and authentication. Domain collections (e.g. chat messages, files lists) are managed using React Query to ensure efficient caching, background updates, and optimistic UI updates. Local component states use standard useState hooks."),
    ("Unit_Testing_Coverage.txt", "Unit Testing Standards and Coverage Enforcements",
     "All Python services must achieve a minimum of 80% test coverage tracked by Coverage.py. Frontend React code requires Jest testing with 75% coverage. Build pipelines check coverage metrics during pull request builds. High-risk modules (auth, payment routers) require 100% test coverage."),
    ("Code_Review_Guidelines.docx", "Engineering Code Review Best Practices",
     "All code contributions require a pull request reviewed by at least one engineer. Reviewers check for: 1. Code correctness. 2. Performance (complexity). 3. Security (no exposed keys). 4. Code styles. PRs must not remain open for more than 48 hours without active discussion. Keep pull requests under 300 lines."),
    ("Feature_Flagging_System.txt", "Feature Flagging and Progressive Rollout Standard",
     "To enable continuous delivery, new features must be wrapped in feature flags using LaunchDarkly. Features are rolled out progressively (1%, 10%, 50%, 100%). Flags are cleaned up and removed from the codebase within 30 days of full release. Feature flags allow decoupling code releases from business launches."),
    ("Logging_and_Tracing_SOP.docx", "Application Logging and Distributed Tracing Setup",
     "All API requests must generate a unique X-Correlation-ID header at the gateway layer. Correlation IDs are passed down to all microservices to enable distributed tracing via OpenTelemetry and Jaeger. Application logs must record severity levels (DEBUG, INFO, WARN, ERROR). Log messages must never expose PII."),
    ("Static_Code_Analysis.txt", "Static Code Analysis and Linting Rules",
     "Python files are linted using Flake8 and formatted using Black. Frontend JS/JSX files are formatted using Prettier and checked via ESLint. Build pipelines automatically reject contributions that violate code linting structures. Code duplicate detection is run monthly, reporting matches > 5%."),
    ("Docker_Compose_Dev.docx", "Docker Compose Local Development Guidelines",
     "For local engineering development, running 'docker-compose up' launches all backing services: MongoDB Atlas local container, Redis server, and active local vector storage. Local code changes are hot-reloaded using volumes mounted to the container instances. Dev configurations are loaded via local .env."),
    ("Server_Uptime_Monitoring.txt", "Infrastructure Performance and Server Uptime",
     "Uptime monitoring runs active ping checks every 60 seconds from three geographical regions (US, Europe, Asia) to our health check endpoint. Latency spikes > 500ms trigger warnings. Consecutive failed pings alert on-call teams via Opsgenie. Status pages are updated automatically showing active availability."),
    ("Mobile_App_Sync.docx", "Mobile App Database Sync Protocols",
     "Mobile apps sync local SQLite databases with the SaaS backend using a change-tracking log mechanism. Synced items utilize a last-write-wins resolver based on UTC timestamps. Large sync payloads are broken into paginated 50-item chunks. Sync networks are compressed using gzip to preserve client data usage."),
    ("Graphql_Schema_Federation.txt", "GraphQL Schema Federation Architecture",
     "Enterprise reporting uses a federated GraphQL gateway built on Apollo Federation. Microservices publish subgraph schemas to Apollo Studio. Schema changes are validated against composition breaking rules. Clients request data through a single gateway endpoint, which optimizes queries across subgraphs.")
]

# Define 30+ Feedback QA Entries to pre-populate database
FEEDBACK_QA_SEED = [
    {"question": "How do partner referrals work?", 
     "answer": "Partners can refer prospective leads through registered agencies and external contractors using the Partner Portal. Successful referrals that sign a 12-month contract earn a 15% commission on the first year's contract value. Commissions are paid out 30 days post invoice clearance.", 
     "tags": "partner, partnerships, agency, external contractors, lead referral, marketing, portal"},
    {"question": "How are database credentials encrypted?", 
     "answer": "Database credentials, API keys, and private certificates must never be stored in plain text or hardcoded in repositories. They must be stored in HashiCorp Vault. Vault access tokens are rotated every 30 days. Databases use IAM-based role authentication where supported, ensuring short-lived access.", 
     "tags": "database, security, encryption, credentials, vault, authentication, rotating keys"},
    {"question": "What is the policy for password rotations?", 
     "answer": "Enterprise passwords must be rotated every 90 days. Passwords must meet a complexity of minimum 14 characters, containing at least one uppercase letter, one lowercase letter, one number, and one special character. Users cannot reuse their previous 5 passwords. Inactive accounts are suspended after 45 days.", 
     "tags": "security, password, rotation, compliance, active directory, complex passwords"},
    {"question": "How do I request admin access?", 
     "answer": "Administrative access requires submitting an Access Request Ticket (ART) signed by your department head. Administrator accounts must use distinct 'admin-' prefixed emails. Access is reviewed quarterly. Temporary elevated privileges (just-in-time access) are limited to 4-hour windows via AWS IAM Identity Center.", 
     "tags": "security, access, admin, provisioning, aws, privileges, art ticket"},
    {"question": "What steps occur in the staging pipeline?", 
     "answer": "The staging environment deployment runs automatically on merges to the 'release' branch. Steps include: 1. Code linting and static analysis (SonarQube). 2. Unit tests execution. 3. Docker image creation. 4. Deployment to EKS staging cluster. 5. E2E integration test suite execution. Staging metrics are logged in Prometheus.", 
     "tags": "ci/cd, pipeline, staging, deployment, docker, testing, prometheus"},
    {"question": "Who approves production hotfixes?", 
     "answer": "Regular production releases require a pull request approved by at least two senior engineers and a signed Change Advisory Board (CAB) ticket. However, hotfixes require immediate VP of Engineering approval, along with automated smoke tests, bypass CAB reviews, and direct hotfix branch deployment.", 
     "tags": "release, production, approvals, hotfixes, cab tickets, smoke tests, deployment"},
    {"question": "Where are docker build logs stored?", 
     "answer": "All Docker container builds export logs in JSON format to stdout. Container host nodes run Vector agents that forward stdout logs directly to Datadog. Docker files must use official base images from verified registries. Multi-stage builds are required to keep final container images under 150MB in size.", 
     "tags": "docker, logs, datadog, vector, containers, dockerfile, storage"},
    {"question": "How does database indexing work?", 
     "answer": "To optimize database read operations, compound indexes must be built on high-cardinality keys. In MongoDB, create compound indexes matching query patterns. Avoid index explosion (maximum 64 indexes per collection). Slow queries taking > 100ms are logged in the slow-query log and reviewed weekly by DBAs.", 
     "tags": "database, indexing, query, optimization, mongodb, dba, slow queries"},
    {"question": "What security protocols are standard?", 
     "answer": "We employ AES-256 encryption at rest and TLS 1.3 in transit. Authentication is managed via Okta SSO with mandatory multi-factor authentication (MFA). Security audits are conducted bi-annually by external Certified Systems Auditors (CISA). Network segmentation is enforced across VPCs.", 
     "tags": "security, standard, encryption, ssl, tls, okta, sso, audits, vpc"},
    {"question": "How do I trigger a blue-green deployment?", 
     "answer": "We execute Blue-Green deployments for stateless microservices. The router (Route53) splits traffic 90/10 during the first 10 minutes of active validation. Once smoke tests succeed on the green pool, 100% traffic is redirected. The previous blue pool remains standby for 60 minutes for immediate rollback ability.", 
     "tags": "deployment, release, blue-green, route53, traffic, rollback, cloud infrastructure"},
    {"question": "What is the data retention policy?", 
     "answer": "Customer database transactional records are retained for 7 years from transaction date to comply with financial regulations. Session logs are retained for 90 days. User account profiles are deleted 30 days after subscription termination. Data scrubbing uses cryptographic deletion (shredding) on block storage devices.", 
     "tags": "data retention, compliance, gdpr, transactional data, shredding, logging"},
    {"question": "What is the API rate limit?", 
     "answer": "The API gateway enforces rate limits to prevent denial of service attacks. Standard tiers are limited to 1,000 requests per hour. Enterprise tiers enjoy 10,000 requests per hour. Rate limits are tracked using a Redis token bucket algorithm. Requests exceeding limits receive an HTTP 429 Too Many Requests status response.", 
     "tags": "api, gateway, rate limiting, throttle, redis, token bucket, http 429"},
    {"question": "How do we isolate customer tenant data?", 
     "answer": "Tenant isolation is maintained using logical schema separation (row-level security) in PostgreSQL databases. Every database query must append a tenant_id filter. Tenant context is populated via JWT tokens at the API middleware layer. Cross-tenant database queries are blocked by strict DB constraint policies.", 
     "tags": "multi-tenant, database, security, isolation, pgsql, row-level security, jwt"},
    {"question": "What is our customer uptime SLA commitment?", 
     "answer": "We commit to a 99.9% uptime SLA for enterprise clients, excluding planned maintenance windows. Planned maintenance is announced 5 business days in advance. Uptime is calculated monthly. Uptime infractions between 99.0% and 99.9% trigger a 10% billing credit; uptime below 99.0% triggers a 25% refund.", 
     "tags": "sla, uptime, customer success, billing, refund, maintenance, metrics"},
    {"question": "What is the employee remote work policy?", 
     "answer": "Employees are permitted to work remotely. All corporate network traffic must route through the corporate tailscale VPN. Working on public Wi-Fi without active VPN is strictly prohibited. Hardware keys (YubiKey) are required for signing into AWS and GitHub. Reimbursed home office setup includes a monitor and chair.", 
     "tags": "hr, remote work, policy, security, vpn, yubikey, equipment"},
    {"question": "What is the expense reimbursement process?", 
     "answer": "Expense reports must be submitted via Expensify within 30 days of purchase. Standard meals are capped at $50 per person. Travel airfares must be booked in economy class. Purchases of software licenses or hardware accessories exceeding $100 require pre-approval from the IT manager. Receipts are mandatory.", 
     "tags": "expense, hr, reimbursement, finance, expensify, travel"},
    {"question": "What happens during employee offboarding?", 
     "answer": "IT schedules automated account revoking at 17:00 local time on the employee's last day. Active accounts revoked: Google Workspace, Slack, AWS, GitHub, Okta, Jira. Company laptops must be returned to the office within 5 business days. Access badges are deactivated immediately.", 
     "tags": "offboarding, hr, security, credentials, assets, laptop, Badge deactivation"},
    {"question": "How do I report code reviews?", 
     "answer": "All code contributions require a pull request reviewed by at least one engineer. Reviewers check for: 1. Code correctness. 2. Performance (complexity). 3. Security (no exposed keys). 4. Code styles. PRs must not remain open for more than 48 hours without active discussion. Keep pull requests under 300 lines.", 
     "tags": "engineering, git, github, code review, pr guidelines, standard"},
    {"question": "How are application logs correlation IDs managed?", 
     "answer": "All API requests must generate a unique X-Correlation-ID header at the gateway layer. Correlation IDs are passed down to all microservices to enable distributed tracing via OpenTelemetry and Jaeger. Application logs must record severity levels (DEBUG, INFO, WARN, ERROR). Log messages must never expose PII.", 
     "tags": "logs, telemetry, correlation id, gateway, opentelemetry, tracing, json logs"},
    {"question": "What is the Redis cache eviction policy?", 
     "answer": "Redis cache instances are configured with maxmemory-policy volatile-lru. All database cache entries must have a designated Time-To-Live (TTL). User sessions have a TTL of 2 hours. Config parameters and tenant settings have a TTL of 24 hours. Cache key formatting uses colon spacing: tenant:{id}:user:{id}.", 
     "tags": "redis, cache, eviction, ttl, lru, sessions, storage"},
    {"question": "How does the mobile app database sync?", 
     "answer": "Mobile apps sync local SQLite databases with the SaaS backend using a change-tracking log mechanism. Synced items utilize a last-write-wins resolver based on UTC timestamps. Large sync payloads are broken into paginated 50-item chunks. Sync networks are compressed using gzip to preserve client data usage.", 
     "tags": "mobile, sync, database, sqlite, compression, last-write-wins"},
    {"question": "What happens when chaos testing runs?", 
     "answer": "We run automated chaos testing weekly using Chaos Mesh in the staging environment. Tests inject latency into DB queries, kill pods randomly, and simulate network partitions. Microservices must demonstrate graceful degradation, responding with cached data or custom fallbacks. Target recovery time objective (RTO) is 15 seconds.", 
     "tags": "chaos engineering, testing, latency, failover, resiliency, kubernetes, rto"},
    {"question": "How are Terraform state files stored?", 
     "answer": "Terraform state files must be stored in a centralized, encrypted S3 bucket. State locking is enforced using a DynamoDB table. State buckets have Versioning enabled. Manual state editing is strictly forbidden. Changes to infrastructure must be committed to git and executed via GitHub Actions CI pipelines.", 
     "tags": "terraform, infrastructure, s3, dynamodb, state lock, devops"},
    {"question": "What is the SLA response time for Tier 3 tickets?", 
     "answer": "Support tickets are categorized: Tier 1: General inquiries (response in 24 hours). Tier 2: Feature bug reports (response in 4 hours). Tier 3: Outage / Security incident (response in 15 minutes). Tier 3 tickets immediately trigger PagerDuty alerts to the on-call engineering team. Support uses Zendesk integrations.", 
     "tags": "support, ticket, sla, response time, Zendesk, PagerDuty, on-call"},
    {"question": "How is Docker container security scanned?", 
     "answer": "Trivy scans are integrated into the GitHub Actions workflow. Pull requests fail if a new container image contains vulnerabilities classified as CRITICAL or HIGH with an available patch. Base images must run as non-root users (UID 10001). We enforce distroless base images for production deployments.", 
     "tags": "security, docker, container, scanning, trivy, vulnerabilities, non-root"},
    {"question": "How is database backup validated?", 
     "answer": "Automated daily database backups are stored in AWS S3 with Glacier vault locks. Every Sunday at 01:00 UTC, an automated restore test launches a temporary DB instance in an isolated subnet, restores the backup, and runs verification SQL tests. Restore status is sent to Slack. Backup failure immediately alerts on-call DBAs.", 
     "tags": "database, backup, validation, s3, slack, alert, disaster recovery"},
    {"question": "What are the rules for feature flagging?", 
     "answer": "To enable continuous delivery, new features must be wrapped in feature flags using LaunchDarkly. Features are rolled out progressively (1%, 10%, 50%, 100%). Flags are cleaned up and removed from the codebase within 30 days of full release. Feature flags allow decoupling code releases from business launches.", 
     "tags": "feature flag, launchdarkly, engineering, continuous delivery, release gating"},
    {"question": "Who conducts the security audits?", 
     "answer": "Security audits are conducted bi-annually by external Certified Information Systems Auditors (CISA). These audits cover penetration testing, compliance validation (SOC2, ISO 27001), and credential checks. Incident logs, Vault policies, and Okta configurations are fully analyzed.", 
     "tags": "security, audits, cisa, compliance, soc2, penetration testing"},
    {"question": "What is the policy for inactive accounts?", 
     "answer": "Inactive user accounts are automatically suspended after 45 days of idle state. System administrators review suspended accounts monthly, and permanent deletion of accounts (scrubbing database keys) is triggered after 180 days of continuous idle state to maintain database cleanliness.", 
     "tags": "security, compliance, inactive accounts, suspension, scrubbing"},
    {"question": "How do we implement GraphQL Schema Federation?", 
     "answer": "Enterprise reporting uses a federated GraphQL gateway built on Apollo Federation. Microservices publish subgraph schemas to Apollo Studio. Schema changes are validated against composition breaking rules. Clients request data through a single gateway endpoint, which optimizes queries across subgraphs.", 
     "tags": "graphql, federation, apollo, gateway, microservices, schemas, design"}
]

async def seed_mongodb(db):
    """Seed the database with default Feedback Q&A pairs."""
    print("Pre-populating feedback_qa database collection...")
    # Clear existing entries
    await db.feedback_qa.delete_many({})
    
    now = datetime.utcnow()
    seeding_entries = []
    for item in FEEDBACK_QA_SEED:
        seeding_entries.append({
            "question": item["question"],
            "answer": item["answer"],
            "tags": item["tags"],
            "createdAt": now,
            "updatedAt": now
        })
        
    result = await db.feedback_qa.insert_many(seeding_entries)
    print(f"Successfully inserted {len(result.inserted_ids)} Feedback Q&A pairs into MongoDB.")

def create_physical_files():
    """Create 52 rich documents locally under sample_documents/ directory."""
    print(f"Creating 52 rich sample files in '{DOCS_DIR}'...")
    
    # Empty existing directory to regenerate
    for f in os.listdir(DOCS_DIR):
        file_p = os.path.join(DOCS_DIR, f)
        if os.path.isfile(file_p):
            os.remove(file_p)
            
    for filename, title, content in DOCUMENT_TOPICS:
        file_path = os.path.join(DOCS_DIR, filename)
        ext = os.path.splitext(filename)[1].lower()
        
        full_content = f"{title}\n{'=' * len(title)}\nDocument Reference: {filename}\nDate Created: 2026-05-29\n\n{content}"
        
        if ext == ".txt":
            # Write plain text
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(full_content)
        elif ext == ".docx":
            # Write a word document using python-docx
            doc = docx.Document()
            doc.add_heading(title, 0)
            doc.add_paragraph(f"Document Reference: {filename}")
            doc.add_paragraph("Date Created: 2026-05-29")
            doc.add_paragraph(content)
            doc.save(file_path)
        elif ext == ".pdf":
            # Write a minimal valid PDF document
            lines = [title, f"Document Reference: {filename}", "Date Created: 2026-05-29", "", content]
            stream_parts = ["BT", "/F1 12 Tf", "50 800 Td", "14 TL"]
            for line in lines:
                # Basic escape for PDF strings
                escaped_line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                stream_parts.append(f"({escaped_line}) Tj")
                stream_parts.append("T*")
            stream_parts.append("ET")
            stream_content = "\n".join(stream_parts).encode('utf-8')
            
            objects = [
                b"<< /Type /Catalog /Pages 2 0 R >>",
                b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
                b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
                f"<< /Length {len(stream_content)} >>\nstream\n".encode('utf-8') + stream_content + b"\nendstream"
            ]
            
            with open(file_path, "wb") as f:
                f.write(b"%PDF-1.4\n")
                offsets = []
                for i, obj in enumerate(objects):
                    offsets.append(f.tell())
                    f.write(f"{i+1} 0 obj\n".encode('utf-8'))
                    f.write(obj)
                    f.write(b"\nendobj\n")
                    
                xref_start = f.tell()
                f.write(b"xref\n")
                f.write(f"0 {len(objects) + 1}\n".encode('utf-8'))
                f.write(b"0000000000 65535 f \n")
                for offset in offsets:
                    f.write(f"{offset:010d} 00000 n \n".encode('utf-8'))
                    
                f.write(b"trailer\n")
                f.write(f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode('utf-8'))
                f.write(b"startxref\n")
                f.write(f"{xref_start}\n".encode('utf-8'))
                f.write(b"%%EOF\n")
            
    print(f"Generated {len(DOCUMENT_TOPICS)} local files.")

async def index_documents_in_rag(db):
    """Run RAG ingestion pipeline over all generated sample documents."""
    print("Starting vector pipeline embedding on the 52 sample documents...")
    # Clear existing document references in MongoDB
    await db.documents.delete_many({})
    
    # Reset vector store manager memory
    vector_store.chunks = []
    vector_store.metadata = []
    vector_store.embeddings = None
    vector_store.faiss_index = None
    if os.path.exists(vector_store.metadata_file):
        os.remove(vector_store.metadata_file)
    if os.path.exists(vector_store.index_file):
        os.remove(vector_store.index_file)
        
    for filename, title, content in DOCUMENT_TOPICS:
        file_path = os.path.join(DOCS_DIR, filename)
        with open(file_path, "rb") as f:
            file_bytes = f.read()
            
        # Parse, chunk, and embed
        text_content = TextProcessor.extract_text(filename, file_bytes)
        chunks = TextProcessor.split_text_into_chunks(text_content, chunk_size=800, chunk_overlap=150)
        chunks_count = len(chunks)
        
        if chunks_count > 0:
            await vector_store.add_documents(filename, chunks, db)
            
        # Save metadata to MongoDB
        doc_entry = {
            "filename": filename,
            "filetype": os.path.splitext(filename)[1].lower().replace(".", ""),
            "uploadDate": datetime.utcnow(),
            "documentPath": file_path,
            "chunksCount": chunks_count
        }
        await db.documents.insert_one(doc_entry)
        print(f"Indexed: {filename} ({chunks_count} chunks)")
        
    await vector_store.save_to_db(db)
    print("Vector storage successfully populated and synchronized with MongoDB!")

async def main():
    print("=== Domain-Specific RAG Database Seeder ===")
    
    # 1. Establish connection to local database
    try:
        from motor.motor_asyncio import AsyncIOMotorClient
        client = AsyncIOMotorClient(settings.MONGODB_URI)
        db = client[settings.DATABASE_NAME]
        
        # Verify connection
        await client.admin.command('ping')
        print("Connected to MongoDB successfully.")
    except Exception as e:
        print(f"ERROR: Could not connect to MongoDB at {settings.MONGODB_URI}.")
        print("Please ensure MongoDB is running or configure the connection URI in backend/.env")
        sys.exit(1)

    # 2. Seed Feedback Database Collection
    await seed_mongodb(db)
    
    # 3. Generate 52 local sample documents
    create_physical_files()
    
    # 4. Process and Index into Vector DB
    await index_documents_in_rag(db)
    
    print("\n===========================================")
    print("Seeding completed successfully!")
    print("The system is now fully pre-loaded with:")
    print("  - 52 Rich operational documents (fully indexed and embedded)")
    print("  - 30 High-quality searchable feedback pairs")
    print("You are ready to launch uvicorn server and run chatbot queries immediately!")
    print("===========================================")

if __name__ == "__main__":
    asyncio.run(main())
